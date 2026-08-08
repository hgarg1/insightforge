"""Excel loading, profiling, transformations, and export helpers."""

from __future__ import annotations

import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import polars as pl


def discover_sheets(path: Path) -> list[str]:
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    except Exception as exc:  # pragma: no cover - dependency-specific fallback
        raise ValueError(f"Unable to inspect workbook sheets: {exc}") from exc


def read_workbook(path: Path, sheet_name: str | None = None) -> pl.DataFrame:
    """Load one sheet with a fast native reader and a reliable fallback."""

    selected = sheet_name or discover_sheets(path)[0]
    try:
        return pl.read_excel(path, sheet_name=selected, engine="calamine")
    except Exception:
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                if selected not in workbook.sheetnames:
                    raise ValueError(f"Sheet '{selected}' was not found")
                worksheet = workbook[selected]
                rows = list(worksheet.values)
            finally:
                workbook.close()
            if not rows:
                return pl.DataFrame()
            headers = _unique_headers(rows[0])
            return pl.DataFrame(rows[1:], schema=headers, orient="row")
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Unable to read workbook: {exc}") from exc


def _unique_headers(values: Any) -> list[str]:
    seen: dict[str, int] = {}
    headers: list[str] = []
    for index, raw in enumerate(values or []):
        base = str(raw).strip() if raw not in (None, "") else f"column_{index + 1}"
        base = re.sub(r"\s+", " ", base)
        count = seen.get(base, 0)
        seen[base] = count + 1
        headers.append(base if count == 0 else f"{base}_{count + 1}")
    return headers


def jsonable(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if hasattr(value, "item"):
        return jsonable(value.item())
    if isinstance(value, dict):
        return {str(key): jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [jsonable(item) for item in value]
    return value


def records(df: pl.DataFrame, limit: int = 100) -> list[dict[str, Any]]:
    return jsonable(df.head(limit).to_dicts())


def _dtype_name(dtype: Any) -> str:
    return str(dtype).replace("<class '", "").replace("'>", "")


def profile_dataframe(df: pl.DataFrame, preview_limit: int = 100) -> dict[str, Any]:
    rows, columns = df.height, df.width
    duplicate_rows = rows - df.unique().height if rows else 0
    total_cells = max(rows * columns, 1)
    null_cells = sum(int(df[column].null_count()) for column in df.columns)
    missing_ratio = null_cells / total_cells
    duplicate_ratio = duplicate_rows / max(rows, 1)
    quality_score = round(max(0.0, min(100.0, 100 - missing_ratio * 60 - duplicate_ratio * 40)), 1)

    column_profiles: list[dict[str, Any]] = []
    for column in df.columns:
        series = df[column]
        non_null = series.drop_nulls()
        profile: dict[str, Any] = {
            "name": column,
            "dtype": _dtype_name(series.dtype),
            "nulls": int(series.null_count()),
            "null_percent": round(float(series.null_count() / max(rows, 1) * 100), 1),
            "unique": int(non_null.n_unique()) if len(non_null) else 0,
            "sample": jsonable(non_null.head(3).to_list()),
        }
        if series.dtype.is_numeric():
            profile.update(
                {
                    "min": jsonable(non_null.min()) if len(non_null) else None,
                    "max": jsonable(non_null.max()) if len(non_null) else None,
                    "mean": jsonable(non_null.mean()) if len(non_null) else None,
                    "median": jsonable(non_null.median()) if len(non_null) else None,
                }
            )
        column_profiles.append(profile)

    return {
        "rows": rows,
        "columns": columns,
        "null_cells": null_cells,
        "duplicate_rows": duplicate_rows,
        "quality_score": quality_score,
        "column_profiles": column_profiles,
        "preview": records(df, preview_limit),
    }


def column_summary(df: pl.DataFrame) -> list[dict[str, Any]]:
    profile = profile_dataframe(df, preview_limit=0)
    return [
        {
            "name": item["name"],
            "dtype": item["dtype"],
            "nulls": item["nulls"],
            "unique": item["unique"],
        }
        for item in profile["column_profiles"]
    ]


def apply_transform(df: pl.DataFrame, operation: dict[str, Any]) -> pl.DataFrame:
    kind = operation.get("operation")
    column = operation.get("column")
    if kind in {"filter", "sort", "rename", "drop_nulls"} and column and column not in df.columns:
        raise ValueError(f"Column '{column}' was not found")

    if kind == "filter":
        operator = operation.get("operator", "eq")
        value = operation.get("value")
        series = df[column]
        if series.dtype.is_numeric():
            try:
                value = float(value)
            except (TypeError, ValueError) as exc:
                raise ValueError("Numeric filters require a numeric value") from exc
        if operator == "contains":
            return df.filter(pl.col(column).cast(pl.String).str.contains(str(value), literal=True))
        comparisons = {
            "eq": pl.col(column) == value,
            "neq": pl.col(column) != value,
            "gt": pl.col(column) > value,
            "gte": pl.col(column) >= value,
            "lt": pl.col(column) < value,
            "lte": pl.col(column) <= value,
        }
        if operator not in comparisons:
            raise ValueError(f"Unsupported filter operator '{operator}'")
        return df.filter(comparisons[operator])

    if kind == "sort":
        return df.sort(column, descending=bool(operation.get("descending", False)))

    if kind == "drop_nulls":
        return df.drop_nulls(subset=[column] if column else None)

    if kind == "rename":
        new_name = str(operation.get("new_name") or "").strip()
        if not new_name:
            raise ValueError("A new column name is required")
        if new_name in df.columns and new_name != column:
            raise ValueError(f"Column '{new_name}' already exists")
        return df.rename({column: new_name})

    if kind == "group":
        measure = operation.get("measure")
        if not column or column not in df.columns:
            raise ValueError("A grouping column is required")
        if measure and measure not in df.columns:
            raise ValueError(f"Measure column '{measure}' was not found")
        aggregation = operation.get("aggregation", "sum")
        if aggregation == "count" or not measure:
            return df.group_by(column).len(name="count").sort("count", descending=True)
        expression = getattr(pl.col(measure), aggregation)()
        return (
            df.group_by(column)
            .agg(expression.alias(aggregation))
            .sort(aggregation, descending=True)
        )

    raise ValueError(f"Unsupported transformation '{kind}'")


def export_dataframe(df: pl.DataFrame, destination: Path, file_format: str = "xlsx") -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if file_format == "csv":
        df.write_csv(destination)
    else:
        df.write_excel(destination, worksheet="InsightForge")


def data_for_chart(df: pl.DataFrame, spec: dict[str, Any]) -> dict[str, Any]:
    """Create a small JSON chart payload consumed by both Flet clients."""

    chart_type = spec.get("chart_type", "bar")
    x_column = spec.get("x_column") or (df.columns[0] if df.columns else None)
    y_column = spec.get("y_column") or (df.columns[1] if len(df.columns) > 1 else x_column)
    color_column = spec.get("color_column")
    if not x_column or x_column not in df.columns:
        raise ValueError("Choose an x-axis column")
    if chart_type != "histogram" and y_column not in df.columns:
        raise ValueError("Choose a y-axis column")

    if chart_type == "histogram":
        values = [
            float(value)
            for value in df[x_column].cast(pl.Float64, strict=False).drop_nulls().to_list()
        ]
        if not values:
            return {"chart_type": chart_type, "x": [], "y": [], "labels": [], "series": []}
        minimum, maximum = min(values), max(values)
        if minimum == maximum:
            return {
                "chart_type": chart_type,
                "x": [jsonable(minimum)],
                "y": [len(values)],
                "labels": [jsonable(minimum)],
                "series": [],
            }
        bin_count = min(10, max(4, int(math.sqrt(len(values)))))
        width = (maximum - minimum) / bin_count
        counts = [0] * bin_count
        for value in values:
            index = min(int((value - minimum) / width), bin_count - 1)
            counts[index] += 1
        labels = [round(minimum + width * index + width / 2, 2) for index in range(bin_count)]
        return {
            "chart_type": chart_type,
            "x": labels,
            "y": counts,
            "labels": labels,
            "series": [],
        }

    if chart_type == "pie":
        grouped = df.group_by(x_column).len(name="count").sort("count", descending=True).head(12)
        return {
            "chart_type": chart_type,
            "x": jsonable(grouped[x_column].to_list()),
            "y": jsonable(grouped["count"].to_list()),
            "labels": jsonable(grouped[x_column].to_list()),
            "series": [],
        }

    numeric = df[y_column].cast(pl.Float64, strict=False)
    base = df.with_columns(numeric.alias("__value"))
    if spec.get("aggregation") and chart_type in {"bar", "line"}:
        aggregation = spec.get("aggregation", "sum")
        if aggregation == "count":
            grouped = base.group_by(x_column).len(name="__value")
        else:
            grouped = base.group_by(x_column).agg(
                getattr(pl.col("__value"), aggregation)().alias("__value")
            )
        grouped = grouped.drop_nulls(subset=["__value"]).sort(x_column).head(100)
    else:
        grouped = base.drop_nulls(subset=["__value"]).select([x_column, "__value"]).head(100)

    series = []
    if color_column and color_column in grouped.columns:
        for label, part in grouped.partition_by(color_column, as_dict=True).items():
            series.append(
                {
                    "name": str(label),
                    "x": jsonable(part[x_column].to_list()),
                    "y": jsonable(part["__value"].to_list()),
                }
            )
    return {
        "chart_type": chart_type,
        "x": jsonable(grouped[x_column].to_list()),
        "y": jsonable(grouped["__value"].to_list()),
        "labels": jsonable(grouped[x_column].to_list()),
        "series": series,
    }
